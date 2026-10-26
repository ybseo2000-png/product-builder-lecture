"""사용자가 내려받은 나라장터 엑셀을 공개용 JSON으로 변환. 셀 내용은 명령으로 실행하지 않음."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
from openpyxl import load_workbook

HEADERS=['업무구분','업무여부','구분','입찰공고번호','공고명','공고기관','수요기관','게시일시(입찰마감일시)']

def parse_row(row):
    values={k:str(v or '').strip() for k,v in zip(HEADERS,row)}
    if not values['입찰공고번호'] or not values['공고명']:raise ValueError('공고번호 또는 공고명 누락')
    if values['공고기관']!='문화체육관광부':raise ValueError('다른 공고기관 포함. 입력 파일 확인 필요')
    dates=re.fullmatch(r'(\d{4}/\d{2}/\d{2} \d{2}:\d{2})\s*\((\d{4}/\d{2}/\d{2} \d{2}:\d{2})\)',values['게시일시(입찰마감일시)'])
    if not dates:raise ValueError('게시/마감일 형식 오류')
    published,deadline=[datetime.strptime(v,'%Y/%m/%d %H:%M').strftime('%Y-%m-%dT%H:%M:00+09:00') for v in dates.groups()]
    return {'id':values['입찰공고번호'],'title':values['공고명'],'org':values['공고기관'],
            'demandOrg':values['수요기관'],'business':values['업무구분'],'procurement':values['업무여부'],
            'noticeType':values['구분'],'publishedAt':published,'deadline':deadline,
            'year':int(published[:4]),'month':int(published[5:7])}

def build(paths):
    records={};sources=[];duplicates=0
    for path in paths:
        path=Path(path)
        workbook=load_workbook(path,data_only=True,read_only=True)
        try:
            for sheet in workbook:
                rows=sheet.iter_rows(values_only=True)
                header=next(rows,None)
                if header is None:continue
                if list(header)!=HEADERS:raise ValueError('지원하지 않는 엑셀 칼럼')
                for row in rows:
                    if not any(v is not None for v in row):continue
                    item=parse_row(row)
                    old=records.get(item['id'])
                    if old:
                        if old!=item:raise ValueError('동일 공고번호에 서로 다른 데이터가 있습니다')
                        duplicates+=1
                    records[item['id']]=item
        finally:workbook.close()
        sources.append({'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
    items=sorted(records.values(),key=lambda r:(r['publishedAt'],r['id']),reverse=True)
    if not items:raise ValueError('빈 파일은 공개 자료를 덮어쓰지 않습니다')
    return {'status':'ready','source':'사용자 제공 나라장터 다운로드 파일','importedAt':datetime.now(timezone.utc).isoformat(),
            'totalCount':len(items),'duplicatesRemoved':duplicates,'sources':sources,
            'dateMin':min(r['publishedAt'] for r in items)[:10],'dateMax':max(r['publishedAt'] for r in items)[:10],
            'industryVerified':False,'hasOriginalUrls':False,'items':items}

def main():
    parser=argparse.ArgumentParser();parser.add_argument('files',nargs='+');parser.add_argument('--output',default='data/nara.json');args=parser.parse_args()
    data=build(args.files);path=Path(args.output);path.parent.mkdir(parents=True,exist_ok=True)
    temp=path.with_suffix('.json.tmp');temp.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');temp.replace(path)
    print(f'수입 완료: {data["totalCount"]}건, 중복 제거 {data["duplicatesRemoved"]}건, {data["dateMin"]}~{data["dateMax"]}')
if __name__=='__main__':main()
